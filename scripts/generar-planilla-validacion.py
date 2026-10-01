"""Genera docs/validacion-expertos.xlsx: la planilla para que un panel de expertos
evalúe la validez de contenido del banco de preguntas de CuidaPiel.

Uso:
    node scripts/exportar-banco.mjs          # actualiza docs/banco-preguntas.json
    python3 scripts/generar-planilla-validacion.py

Requiere openpyxl. Las fórmulas (I-CVI, kappa modificado, S-CVI/Ave, S-CVI/UA) siguen a
Lynn 1986 (PMID 3640358), Polit y Beck 2006 (PMID 16977646) y Polit, Beck y Owen 2007
(PMID 17654487).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
banco = json.loads((RAIZ / "docs" / "banco-preguntas.json").read_text(encoding="utf-8"))
items = banco["items"]
N_EXP = 6
N = len(items)
ULT = N + 1  # última fila de datos (encabezado en la fila 1)

FUENTE = "Arial"
f_base = Font(name=FUENTE, size=10)
f_negrita = Font(name=FUENTE, size=10, bold=True)
f_titulo = Font(name=FUENTE, size=14, bold=True, color="1F2A6E")
f_encabezado = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
f_nota = Font(name=FUENTE, size=9, italic=True, color="3B4270")
relleno_enc = PatternFill("solid", fgColor="1F2A6E")
relleno_input = PatternFill("solid", fgColor="FFF2CC")
relleno_calc = PatternFill("solid", fgColor="EEF0FB")
borde = Border(bottom=Side(style="thin", color="D0D3E6"))
ajustar = Alignment(wrap_text=True, vertical="top")
centro = Alignment(horizontal="center", vertical="top")

FORMA = {"A": "A", "B": "B", "D": "D"}
NOMBRE_FORMA = {"A": "A · ingreso", "B": "B · alta", "D": "D · desafío diario"}


def encabezados(ws, titulos, anchos):
    for i, (t, w) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.font = f_encabezado
        c.fill = relleno_enc
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "B2"


wb = Workbook()

# ---------------------------------------------------------------- Instrucciones
ws = wb.active
ws.title = "Instrucciones"
ws.column_dimensions["A"].width = 30
for col in "BCDEFGHIJ":
    ws.column_dimensions[col].width = 12
lineas = [
    ("CuidaPiel · Validación de contenido del banco de preguntas", f_titulo),
    (f"Versión del banco: {banco['version']} · {N} ítems (12 forma A de ingreso, 12 forma B de alta, 18 del desafío diario)", f_base),
    ("", f_base),
    ("Qué hacer", f_negrita),
    ("1. Cada experto o experta tiene una columna (E1 a E6) en las hojas «Relevancia» y «Claridad». Complete solo las celdas amarillas.", f_base),
    ("2. Relevancia: ¿el ítem mide un cuidado importante para prevenir lesiones por presión en casa? 1 = no relevante · 2 = algo relevante · 3 = bastante relevante · 4 = muy relevante.", f_base),
    ("3. Claridad: ¿lo entiende un paciente chileno con poca escolaridad? 1 = nada claro · 2 = poco claro · 3 = claro · 4 = muy claro. Anote sugerencias de redacción en la columna final.", f_base),
    ("4. Si hay menos de 6 expertos, deje vacías las columnas que sobran: las fórmulas cuentan solo las celdas con nota.", f_base),
    ("5. En «Relevancia», la columna «Decisión» (amarilla) es para el equipo investigador: Mantener, Modificar o Eliminar.", f_base),
    ("", f_base),
    ("Cómo se calcula", f_negrita),
    ("I-CVI = expertos que puntúan 3 o 4 ÷ expertos que evaluaron el ítem (Lynn 1986; Polit y Beck 2006).", f_base),
    ("Kappa modificado κ* = (I-CVI − Pc) ÷ (1 − Pc), con Pc = COMBIN(N; A) × 0,5^N, la probabilidad de acuerdo por azar (Polit, Beck y Owen 2007).", f_base),
    ("Interpretación de κ*: > 0,74 excelente · 0,60–0,74 bueno · 0,40–0,59 regular · < 0,40 revisar (Polit, Beck y Owen 2007).", f_base),
    ("S-CVI/Ave = promedio de los I-CVI; S-CVI/UA = proporción de ítems con I-CVI = 1 (Polit y Beck 2006). Ver hoja «Resumen».", f_base),
    ("", f_base),
    ("Ejemplo de una fila completa (no forma parte del banco)", f_negrita),
]
for r, (txt, fnt) in enumerate(lineas, start=1):
    c = ws.cell(row=r, column=1, value=txt)
    c.font = fnt
ej = len(lineas) + 1
for i, t in enumerate(["Ítem de ejemplo", "E1", "E2", "E3", "E4", "E5", "E6", "N", "I-CVI", "κ*"]):
    c = ws.cell(row=ej, column=1 + i, value=t)
    c.font = f_encabezado
    c.fill = relleno_enc
    c.alignment = centro
ws.cell(row=ej + 1, column=1, value="EJEMPLO").font = f_base
for i, v in enumerate([4, 4, 3, 4, 3, 2]):
    c = ws.cell(row=ej + 1, column=2 + i, value=v)
    c.font = Font(name=FUENTE, size=10, color="0000FF")
    c.fill = relleno_input
    c.alignment = centro
r = ej + 1
ws.cell(row=r, column=8, value=f"=COUNT(B{r}:G{r})")
ws.cell(row=r, column=9, value=f"=IF(H{r}=0,\"\",COUNTIF(B{r}:G{r},\">=3\")/H{r})")
ws.cell(row=r, column=10, value=f"=IF(H{r}=0,\"\",(I{r}-COMBIN(H{r},COUNTIF(B{r}:G{r},\">=3\"))*0.5^H{r})/(1-COMBIN(H{r},COUNTIF(B{r}:G{r},\">=3\"))*0.5^H{r}))")
for col in (8, 9, 10):
    c = ws.cell(row=r, column=col)
    c.font = f_base
    c.fill = relleno_calc
    c.alignment = centro
ws.cell(row=r, column=9).number_format = "0.00"
ws.cell(row=r, column=10).number_format = "0.00"
ws.cell(row=r + 2, column=1, value="Leyenda: celdas amarillas = para completar (números azules) · celdas lilas = cálculo automático; no las edite.").font = f_nota
ws.cell(row=r + 3, column=1, value="Referencias: Lynn MR. Nurs Res. 1986;35(6):382-5 (PMID 3640358). Polit DF, Beck CT. Res Nurs Health. 2006;29(5):489-97 (PMID 16977646). Polit DF, Beck CT, Owen SV. Res Nurs Health. 2007;30(4):459-67 (PMID 17654487).").font = f_nota

# ---------------------------------------------------------------- Relevancia
rel = wb.create_sheet("Relevancia")
exp_cols = [f"E{i}" for i in range(1, N_EXP + 1)]
titulos = ["ID", "Forma", "Dominio", "Competencia", "Enunciado", "Respuesta correcta", *exp_cols,
           "N expertos", "Acuerdo (3 o 4)", "I-CVI", "Pc (azar)", "κ*", "Interpretación", "Decisión", "Comentarios"]
anchos = [8, 7, 18, 26, 46, 34, *[5] * N_EXP, 9, 9, 8, 9, 8, 13, 13, 30]
encabezados(rel, titulos, anchos)
c_e1 = 7
c_eN = c_e1 + N_EXP - 1
L1, LN = get_column_letter(c_e1), get_column_letter(c_eN)
cN, cA, cI, cP, cK, cInt, cDec, cCom = (get_column_letter(c_eN + k) for k in range(1, 9))
for i, it in enumerate(items, start=2):
    vals = [it["id"], FORMA[it["forma"]], banco["dominios"][it["dominio"]],
            f"{it['comp']} · {banco['competencias'][it['comp']]['nombre']}", it["t"], it["op"][it["ok"]]]
    for j, v in enumerate(vals, start=1):
        c = rel.cell(row=i, column=j, value=v)
        c.font = f_base
        c.alignment = ajustar
        c.border = borde
    for j in range(c_e1, c_eN + 1):
        c = rel.cell(row=i, column=j)
        c.fill = relleno_input
        c.font = Font(name=FUENTE, size=10, color="0000FF")
        c.alignment = centro
        c.border = borde
    formulas = {
        cN: f"=COUNT({L1}{i}:{LN}{i})",
        cA: f"=COUNTIF({L1}{i}:{LN}{i},\">=3\")",
        cI: f"=IF({cN}{i}=0,\"\",{cA}{i}/{cN}{i})",
        cP: f"=IF({cN}{i}=0,\"\",COMBIN({cN}{i},{cA}{i})*0.5^{cN}{i})",
        cK: f"=IF({cN}{i}=0,\"\",({cI}{i}-{cP}{i})/(1-{cP}{i}))",
        cInt: f"=IF({cK}{i}=\"\",\"\",IF({cK}{i}>0.74,\"Excelente\",IF({cK}{i}>=0.6,\"Bueno\",IF({cK}{i}>=0.4,\"Regular\",\"Revisar\"))))",
    }
    for col, fml in formulas.items():
        c = rel[f"{col}{i}"]
        c.value = fml
        c.font = f_base
        c.fill = relleno_calc
        c.alignment = centro
        c.border = borde
    rel[f"{cI}{i}"].number_format = "0.00"
    rel[f"{cP}{i}"].number_format = "0.000"
    rel[f"{cK}{i}"].number_format = "0.00"
    for col in (cDec, cCom):
        c = rel[f"{col}{i}"]
        c.fill = relleno_input
        c.font = Font(name=FUENTE, size=10, color="0000FF")
        c.alignment = ajustar
        c.border = borde
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True,
                    errorTitle="Valor no válido", error="Use un número entero de 1 a 4.")
rel.add_data_validation(dv)
dv.add(f"{L1}2:{LN}{ULT}")
dv_dec = DataValidation(type="list", formula1='"Mantener,Modificar,Eliminar"', allow_blank=True)
rel.add_data_validation(dv_dec)
dv_dec.add(f"{cDec}2:{cDec}{ULT}")
rel[f"{cI}1"].comment = Comment("I-CVI: proporción de expertos que puntúan 3 o 4 (Lynn 1986; Polit y Beck 2006).", "CuidaPiel")
rel[f"{cK}1"].comment = Comment("Kappa modificado: corrige el I-CVI por acuerdo al azar (Polit, Beck y Owen 2007).", "CuidaPiel")

# ---------------------------------------------------------------- Claridad
cla = wb.create_sheet("Claridad")
titulos = ["ID", "Forma", "Enunciado", "Alternativas", *exp_cols, "N expertos", "Claro (3 o 4)", "Índice de claridad", "Sugerencias de redacción"]
anchos = [8, 7, 46, 52, *[5] * N_EXP, 9, 9, 10, 40]
encabezados(cla, titulos, anchos)
k1 = 5
kN = k1 + N_EXP - 1
K1, KN = get_column_letter(k1), get_column_letter(kN)
qN, qA, qI, qS = (get_column_letter(kN + k) for k in range(1, 5))
for i, it in enumerate(items, start=2):
    alts = " / ".join(f"{'ABC'[k]}) {o}" for k, o in enumerate(it["op"]))
    for j, v in enumerate([it["id"], FORMA[it["forma"]], it["t"], alts], start=1):
        c = cla.cell(row=i, column=j, value=v)
        c.font = f_base
        c.alignment = ajustar
        c.border = borde
    for j in range(k1, kN + 1):
        c = cla.cell(row=i, column=j)
        c.fill = relleno_input
        c.font = Font(name=FUENTE, size=10, color="0000FF")
        c.alignment = centro
        c.border = borde
    for col, fml in {qN: f"=COUNT({K1}{i}:{KN}{i})", qA: f"=COUNTIF({K1}{i}:{KN}{i},\">=3\")",
                     qI: f"=IF({qN}{i}=0,\"\",{qA}{i}/{qN}{i})"}.items():
        c = cla[f"{col}{i}"]
        c.value = fml
        c.font = f_base
        c.fill = relleno_calc
        c.alignment = centro
        c.border = borde
    cla[f"{qI}{i}"].number_format = "0.00"
    c = cla[f"{qS}{i}"]
    c.fill = relleno_input
    c.font = Font(name=FUENTE, size=10, color="0000FF")
    c.alignment = ajustar
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True,
                     errorTitle="Valor no válido", error="Use un número entero de 1 a 4.")
cla.add_data_validation(dv2)
dv2.add(f"{K1}2:{KN}{ULT}")

# ---------------------------------------------------------------- Formas paralelas
par = wb.create_sheet("Formas paralelas")
encabezados(par, ["Competencia", "Dominio", "Ítem A (ingreso)", "I-CVI A", "Ítem B (alta)", "I-CVI B", "Ambos ≥ umbral"],
            [42, 24, 14, 10, 14, 10, 14])
fila = 2
for comp, info in banco["competencias"].items():
    ia = next(it["id"] for it in items if it["comp"] == comp and it["forma"] == "A")
    ib = next(it["id"] for it in items if it["comp"] == comp and it["forma"] == "B")
    vals = [f"{comp} · {info['nombre']}", banco["dominios"][info["dominio"]], ia,
            f"=IFERROR(INDEX(Relevancia!${cI}$2:${cI}${ULT},MATCH(C{fila},Relevancia!$A$2:$A${ULT},0)),\"\")",
            ib,
            f"=IFERROR(INDEX(Relevancia!${cI}$2:${cI}${ULT},MATCH(E{fila},Relevancia!$A$2:$A${ULT},0)),\"\")",
            f"=IF(OR(D{fila}=\"\",F{fila}=\"\"),\"\",IF(AND(D{fila}>=Resumen!$B$3,F{fila}>=Resumen!$B$3),\"Sí\",\"No\"))"]
    for j, v in enumerate(vals, start=1):
        c = par.cell(row=fila, column=j, value=v)
        c.font = Font(name=FUENTE, size=10, color="008000") if j in (4, 6) else f_base
        c.alignment = ajustar if j <= 2 else centro
        c.border = borde
    par.cell(row=fila, column=4).number_format = "0.00"
    par.cell(row=fila, column=6).number_format = "0.00"
    fila += 1

# ---------------------------------------------------------------- Resumen
res = wb.create_sheet("Resumen")
res.column_dimensions["A"].width = 30
for col in "BCDEFG":
    res.column_dimensions[col].width = 16
res["A1"] = "Resumen de validez de contenido"
res["A1"].font = f_titulo
res["A3"] = "Umbral I-CVI aceptable"
res["B3"] = 0.78
res["A4"] = "Umbral S-CVI/Ave aceptable"
res["B4"] = 0.9
for ref in ("A3", "A4"):
    res[ref].font = f_base
for ref in ("B3", "B4"):
    res[ref].font = Font(name=FUENTE, size=10, color="0000FF")
    res[ref].fill = relleno_input
    res[ref].number_format = "0.00"
res["C3"] = "Polit y Beck 2006 (PMID 16977646): I-CVI ≥ 0,78 con 6 o más expertos (Lynn 1986)."
res["C4"] = "Polit y Beck 2006 (PMID 16977646): S-CVI/Ave ≥ 0,90."
res["C3"].font = f_nota
res["C4"].font = f_nota

hdr = ["Conjunto", "Ítems", "Ítems evaluados", "S-CVI/Ave", "S-CVI/UA", "Ítems bajo umbral", "Cumple S-CVI/Ave"]
for j, t in enumerate(hdr, start=1):
    c = res.cell(row=6, column=j, value=t)
    c.font = f_encabezado
    c.fill = relleno_enc
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
rB = f"Relevancia!$B$2:$B${ULT}"
rN = f"Relevancia!${cN}$2:${cN}${ULT}"
rI = f"Relevancia!${cI}$2:${cI}${ULT}"
rC = f"Relevancia!$C$2:$C${ULT}"
fila = 7
for clave in ("A", "B", "D"):
    crit = f"\"{clave}\""
    vals = [NOMBRE_FORMA[clave],
            f"=COUNTIF({rB},{crit})",
            f"=COUNTIFS({rB},{crit},{rN},\">0\")",
            f"=IFERROR(AVERAGEIFS({rI},{rB},{crit}),\"\")",
            f"=IF(C{fila}=0,\"\",COUNTIFS({rB},{crit},{rI},1)/C{fila})",
            f"=IF(C{fila}=0,\"\",COUNTIFS({rB},{crit},{rI},\"<\"&$B$3))",
            f"=IF(D{fila}=\"\",\"\",IF(D{fila}>=$B$4,\"Sí\",\"No\"))"]
    for j, v in enumerate(vals, start=1):
        c = res.cell(row=fila, column=j, value=v)
        c.font = f_base
        c.alignment = centro if j > 1 else ajustar
    fila += 1
vals = ["Banco completo", f"=COUNTA(Relevancia!$A$2:$A${ULT})", f"=COUNTIF({rN},\">0\")",
        f"=IFERROR(AVERAGE({rI}),\"\")", f"=IF(C{fila}=0,\"\",COUNTIF({rI},1)/C{fila})",
        f"=IF(C{fila}=0,\"\",COUNTIF({rI},\"<\"&$B$3))", f"=IF(D{fila}=\"\",\"\",IF(D{fila}>=$B$4,\"Sí\",\"No\"))"]
for j, v in enumerate(vals, start=1):
    c = res.cell(row=fila, column=j, value=v)
    c.font = f_negrita
    c.alignment = centro if j > 1 else ajustar
for r in range(7, fila + 1):
    res[f"D{r}"].number_format = "0.00"
    res[f"E{r}"].number_format = "0.00"

fila += 2
for j, t in enumerate(["Dominio", "Ítems", "S-CVI/Ave"], start=1):
    c = res.cell(row=fila, column=j, value=t)
    c.font = f_encabezado
    c.fill = relleno_enc
    c.alignment = Alignment(horizontal="center")
for nombre in banco["dominios"].values():
    fila += 1
    res.cell(row=fila, column=1, value=nombre).font = f_base
    res.cell(row=fila, column=2, value=f"=COUNTIF({rC},A{fila})").font = f_base
    c = res.cell(row=fila, column=3, value=f"=IFERROR(AVERAGEIFS({rI},{rC},A{fila}),\"\")")
    c.font = f_base
    c.number_format = "0.00"
res.cell(row=fila + 2, column=1, value="Las celdas quedan vacías hasta que haya notas de expertos en la hoja «Relevancia».").font = f_nota

wb.move_sheet("Resumen", offset=-3)
# Sin valores en caché: Excel y LibreOffice recalculan todo al abrir el archivo.
wb.calculation.fullCalcOnLoad = True
destino = RAIZ / "docs" / "validacion-expertos.xlsx"
wb.save(destino)
print(f"Planilla escrita en {destino.relative_to(RAIZ)} ({N} ítems, {N_EXP} columnas de expertos)")
